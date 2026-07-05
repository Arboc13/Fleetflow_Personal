import datetime as dt
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.fuel import FuelTransaction
from app.models.user import Role

V1 = {
    "plate": "B-100-AAA", "vin": "WVWZZZ1JZXW111111", "make": "Dacia", "model": "Logan",
    "year": 2021, "fuel_type": "petrol", "tank_capacity_l": 50, "current_km": 20000,
}
V2 = {
    "plate": "B-200-BBB", "vin": "WVWZZZ1JZXW222222", "make": "Ford", "model": "Focus",
    "year": 2022, "fuel_type": "diesel", "tank_capacity_l": 60, "current_km": 0,
}
D1 = {
    "email": "sofer@example.com", "full_name": "Sofer Unu", "password": "drivepass1",
    "cnp": "1900101223344", "license_number": "L-1", "license_category": "B",
    "license_expiry": "2030-01-01",
}

PERIOD = {"date_from": "2026-06-01", "date_to": "2026-06-30"}  # 30 days


def _seed(client: TestClient, admin: dict, db: Session) -> tuple[int, int]:
    """Hand-computed dataset: V1 costs 350 fuel + 500 service + 100 docs = 950,
    drives 500 km (cost/km 1.9), allocated 15 of 30 days (50%). V2 is idle."""
    v1 = client.post("/api/vehicles", json=V1, headers=admin).json()["id"]
    v2 = client.post("/api/vehicles", json=V2, headers=admin).json()["id"]
    d1 = client.post("/api/drivers", json=D1, headers=admin).json()["id"]

    client.post(
        "/api/allocations",
        json={"vehicle_id": v1, "driver_id": d1, "start_at": "2026-06-01T00:00:00",
              "end_at": "2026-06-16T00:00:00", "type": "trip"},
        headers=admin,
    )
    tid = client.post(
        "/api/trip-sheets",
        json={"vehicle_id": v1, "driver_id": d1,
              "departure_at": "2026-06-02T08:00:00", "start_km": 20000},
        headers=admin,
    ).json()["id"]
    client.post(
        f"/api/trip-sheets/{tid}/close",
        json={"arrival_at": "2026-06-02T17:00:00", "end_km": 20500},
        headers=admin,
    )
    client.post(
        "/api/service-records",
        json={"vehicle_id": v1, "date": "2026-06-10", "km_at_service": 20100,
              "work_description": "Oil change", "cost": 500},
        headers=admin,
    )
    client.post(
        "/api/documents",
        json={"vehicle_id": v1, "type": "RCA", "series_number": "RO-1", "cost": 100,
              "issue_date": "2026-06-05", "expiry_date": "2027-06-05"},
        headers=admin,
    )
    db.add(  # in period: 50 l x 7.00 = 350
        FuelTransaction(
            vehicle_id=v1, occurred_at=dt.datetime(2026, 6, 3, 10, 0), liters=50, price=7
        )
    )
    db.add(  # outside the period: must be excluded
        FuelTransaction(
            vehicle_id=v1, occurred_at=dt.datetime(2026, 7, 2, 10, 0), liters=100, price=10
        )
    )
    db.commit()
    return v1, v2


def test_tco_report_matches_hand_computed_values(
    client: TestClient, auth_headers, db_session: Session
) -> None:
    admin = auth_headers(Role.admin)
    v1, v2 = _seed(client, admin, db_session)

    report = client.get("/api/reports/tco", params=PERIOD, headers=admin).json()
    rows = {r["vehicle_id"]: r for r in report["rows"]}

    r1 = rows[v1]
    assert r1["fuel_cost"] == 350.0
    assert r1["service_cost"] == 500.0
    assert r1["document_cost"] == 100.0
    assert r1["total_cost"] == 950.0
    assert r1["km_driven"] == 500
    assert r1["cost_per_km"] == 1.9
    assert r1["utilization_pct"] == 50.0

    r2 = rows[v2]
    assert r2["total_cost"] == 0.0
    assert r2["km_driven"] == 0
    assert r2["cost_per_km"] is None
    assert r2["utilization_pct"] == 0.0

    fleet = report["fleet"]
    assert fleet["total_cost"] == 950.0
    assert fleet["km_driven"] == 500
    assert fleet["cost_per_km"] == 1.9
    assert fleet["utilization_pct"] == 25.0  # mean of 50% and 0%


def test_tco_vehicle_filter_and_date_validation(
    client: TestClient, auth_headers, db_session: Session
) -> None:
    admin = auth_headers(Role.admin)
    v1, _ = _seed(client, admin, db_session)

    report = client.get(
        "/api/reports/tco", params={**PERIOD, "vehicle_id": v1}, headers=admin
    ).json()
    assert len(report["rows"]) == 1
    assert report["rows"][0]["vehicle_id"] == v1

    assert client.get(
        "/api/reports/tco",
        params={"date_from": "2026-06-30", "date_to": "2026-06-01"},
        headers=admin,
    ).status_code == 422


def test_tco_xlsx_export_has_numeric_cells(
    client: TestClient, auth_headers, db_session: Session
) -> None:
    admin = auth_headers(Role.admin)
    _seed(client, admin, db_session)

    resp = client.get(
        "/api/reports/tco/export", params={**PERIOD, "format": "xlsx"}, headers=admin
    )
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]
    assert ".xlsx" in resp.headers["content-disposition"]

    ws = load_workbook(BytesIO(resp.content)).active
    # Row 3 = first vehicle (after title + header). Column G = total cost.
    # Real numbers (not strings), usable in Excel formulas (DoD #4).
    assert ws["G3"].value == 950
    assert isinstance(ws["G3"].value, (int, float))
    assert ws["H3"].value == 500  # km driven
    assert isinstance(ws["H3"].value, (int, float))


def test_tco_pdf_export_and_rbac(
    client: TestClient, auth_headers, db_session: Session
) -> None:
    admin = auth_headers(Role.admin)
    _seed(client, admin, db_session)

    resp = client.get(
        "/api/reports/tco/export", params={**PERIOD, "format": "pdf"}, headers=admin
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert resp.content.startswith(b"%PDF")

    driver = auth_headers(Role.driver, email="d@x.com")
    assert client.get("/api/reports/tco", params=PERIOD, headers=driver).status_code == 403
